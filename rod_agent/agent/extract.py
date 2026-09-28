"""업로드 파일 → Claude 입력 블록.

PDF·이미지는 원본 그대로(문서/이미지 블록) 보내 표·스캔본까지 읽게 하고,
나머지는 텍스트로 추출한다. 구형 .hwp처럼 자동 추출이 불가능한 파일은
추측하지 않고 '수동 확인 필요'로 표시한다(지시서 §6).
"""

from __future__ import annotations

import base64
import csv
import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
               ".gif": "image/gif", ".webp": "image/webp"}
TEXT_TYPES = {".txt", ".md", ".json", ".py", ".js", ".html", ".css", ".log", ".bat", ".yaml", ".yml"}
MAX_TEXT = 150_000
MAX_PDF_BYTES = 30 * 1024 * 1024
MAX_IMAGE_BYTES = 5 * 1024 * 1024
UNSUPPORTED = {".hwp": "구형 한글(.hwp)은 자동 추출 불가 — 한글에서 HWPX 또는 PDF로 저장 후 다시 올려주세요."}


def _pdf_text(path: Path) -> str:
    import pymupdf

    with pymupdf.open(path) as doc:
        return "\n\n".join(f"[p.{i + 1}]\n{page.get_text()}" for i, page in enumerate(doc))


def _docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    out = [p.text for p in doc.paragraphs if p.text.strip()]
    for t_idx, table in enumerate(doc.tables, start=1):
        out.append(f"\n[표 {t_idx}]")
        for row in table.rows:
            out.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
    return "\n".join(out)


def _xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"\n[시트: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append("| " + " | ".join("" if v is None else str(v) for v in row) + " |")
    return "\n".join(out)


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        out.append(f"\n[슬라이드 {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append("| " + " | ".join(c.text for c in row.cells) + " |")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("노트: " + slide.notes_slide.notes_text_frame.text)
    return "\n".join(out)


def _hwpx_text(path: Path) -> str:
    out = []
    with zipfile.ZipFile(path) as zf:
        sections = sorted(n for n in zf.namelist() if re.match(r"Contents/section\d+\.xml$", n))
        for name in sections:
            root = ET.fromstring(zf.read(name))
            for p in root.iter():
                if p.tag.endswith("}p"):
                    text = "".join(t.text or "" for t in p.iter() if t.tag.endswith("}t"))
                    # 표 안 문단은 부모 문단 텍스트에도 포함되므로 중복을 피한다
                    if text.strip() and not any(c.tag.endswith("}tbl") for c in p.iter()):
                        out.append(text)
    return "\n".join(out)


def _csv_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp949"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    rows = csv.reader(io.StringIO(text))
    return "\n".join("| " + " | ".join(r) + " |" for r in rows)


def _plain_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp949"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


EXTRACTORS = {
    ".pdf": _pdf_text, ".docx": _docx_text, ".xlsx": _xlsx_text, ".xlsm": _xlsx_text,
    ".pptx": _pptx_text, ".hwpx": _hwpx_text, ".csv": _csv_text,
}


def analyze(path: Path) -> dict:
    """업로드 직후 1회 호출: 추출 텍스트와 상태를 돌려준다."""
    ext = path.suffix.lower()
    info = {"kind": "text", "text": "", "status": "ok", "note": ""}
    if ext in UNSUPPORTED:
        info.update(kind="unsupported", status="manual", note=UNSUPPORTED[ext])
        return info
    if ext in IMAGE_TYPES:
        if path.stat().st_size > MAX_IMAGE_BYTES:
            info.update(kind="unsupported", status="manual", note="이미지가 5MB를 넘습니다. 크기를 줄여 다시 올려주세요.")
        else:
            info.update(kind="image", note="이미지")
        return info
    try:
        if ext in EXTRACTORS:
            info["text"] = EXTRACTORS[ext](path)
        elif ext in TEXT_TYPES:
            info["text"] = _plain_text(path)
        else:
            info.update(kind="unsupported", status="manual", note=f"{ext} 형식은 자동 추출을 지원하지 않습니다.")
            return info
    except Exception as e:  # noqa: BLE001 - 손상 파일 등은 수동 확인으로 돌린다
        info.update(kind="unsupported", status="manual", note=f"추출 실패: {e}")
        return info
    if ext == ".pdf":
        info["kind"] = "pdf" if path.stat().st_size <= MAX_PDF_BYTES else "text"
        if not info["text"].strip():
            info["note"] = "텍스트 층이 없는 스캔 PDF — 원본 이미지로 판독합니다."
    if len(info["text"]) > MAX_TEXT:
        info["note"] = f"내용이 길어 앞 {MAX_TEXT:,}자만 사용합니다."
    return info


def to_blocks(path: Path, original_name: str, info: dict) -> list[dict]:
    """Claude messages API용 content 블록."""
    header = f"[첨부파일: {original_name}]"
    if info["kind"] == "pdf":
        data = base64.standard_b64encode(path.read_bytes()).decode()
        return [
            {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": data},
             "title": original_name},
        ]
    if info["kind"] == "image":
        data = base64.standard_b64encode(path.read_bytes()).decode()
        return [
            {"type": "text", "text": header},
            {"type": "image", "source": {"type": "base64", "media_type": IMAGE_TYPES[path.suffix.lower()], "data": data}},
        ]
    if info["kind"] == "unsupported":
        return [{"type": "text", "text": f"{header}\n(자동 추출 불가 — 수동 확인 필요: {info['note']}) 이 파일 내용을 추측하지 말 것."}]
    return [{"type": "text", "text": f"{header}\n{info['text'][:MAX_TEXT]}"}]
