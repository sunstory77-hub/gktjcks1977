"""마크다운 → 납품 파일(HWPX/DOCX/PPTX/XLSX) 생성과 검증."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches
from pptx.util import Pt as PptPt

from . import hwpx_writer
from .markdown import Block, extract_meta, parse, plain, split_bold, title_of

FORMATS = ("hwpx", "docx", "pptx", "xlsx")
FONT = "맑은 고딕"
ACCENT = (0x1F, 0x4E, 0x79)


# ---------------------------------------------------------------- 파일명

def _safe(part: str) -> str:
    part = re.sub(r'[\\/:*?"<>|\r\n\t]+', " ", part or "").strip()
    return re.sub(r"\s+", "_", part)[:60]


def build_filename(out_dir: Path, doc_type: str, ext: str, org: str = "") -> Path:
    """[기관명]_[문서종류]_[YYYYMMDD].ext, 같은 이름이 있으면 _v2, _v3 …"""
    stem = "_".join(p for p in (_safe(org), _safe(doc_type) or "문서", date.today().strftime("%Y%m%d")) if p)
    path = out_dir / f"{stem}.{ext}"
    n = 2
    while path.exists():
        path = out_dir / f"{stem}_v{n}.{ext}"
        n += 1
    return path


# ---------------------------------------------------------------- DOCX

def _docx_font(run, size: float | None = None, bold: bool | None = None, color=None) -> None:
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def _docx_text(paragraph, text: str, size: float = 10.5) -> None:
    for chunk, bold in split_bold(text):
        _docx_font(paragraph.add_run(chunk), size=size, bold=bold or None)


def write_docx(blocks: list[Block], path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)

    first = True
    for b in blocks:
        if b.kind == "heading":
            if first and b.level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _docx_font(p.add_run(plain(b.text)), size=20, bold=True, color=ACCENT)
            else:
                p = doc.add_heading(level=min(b.level, 4))
                size = {1: 16, 2: 14, 3: 12}.get(b.level, 11)
                _docx_font(p.add_run(plain(b.text)), size=size, bold=True, color=ACCENT)
            first = False
        elif b.kind == "para":
            _docx_text(doc.add_paragraph(), b.text)
        elif b.kind == "bullet":
            style = "List Number" if b.ordered else "List Bullet"
            if b.level:
                style += f" {min(b.level + 1, 3)}"
            _docx_text(doc.add_paragraph(style=style), b.text)
        elif b.kind == "quote":
            p = doc.add_paragraph(style="Intense Quote")
            _docx_text(p, b.text)
        elif b.kind == "table" and b.rows:
            table = doc.add_table(rows=len(b.rows), cols=len(b.rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(b.rows):
                for c, cell in enumerate(row):
                    para = table.cell(r, c).paragraphs[0]
                    if r == 0:
                        _docx_font(para.add_run(plain(cell)), size=10, bold=True)
                        table.cell(r, c)._tc.get_or_add_tcPr().append(_shade("D9E2F3"))
                    else:
                        _docx_text(para, cell, size=10)
            doc.add_paragraph()
        elif b.kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(b.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif b.kind == "hr":
            doc.add_paragraph()
    doc.save(path)


def _shade(hex_color: str):
    from docx.oxml import OxmlElement

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    return shd


# ---------------------------------------------------------------- PPTX

_MAX_LINES = 9


def _slides_from_blocks(blocks: list[Block]) -> tuple[str, str, list[dict]]:
    """H1 → 표지, H2(또는 ---) → 새 슬라이드, '노트:'/인용 → 발표자 노트."""
    title, subtitle = title_of(blocks, "강의자료"), ""
    slides: list[dict] = []
    cur: dict | None = None
    seen_title = False

    def new(t: str) -> dict:
        s = {"title": t, "items": [], "notes": []}
        slides.append(s)
        return s

    for b in blocks:
        if b.kind == "heading" and b.level == 1 and not seen_title:
            seen_title = True
            continue
        if b.kind == "heading" and b.level <= 2:
            cur = new(plain(b.text))
            continue
        if b.kind == "hr":
            cur = None
            continue
        if not seen_title and not slides and b.kind == "para" and not subtitle:
            subtitle = plain(b.text)
            continue
        if cur is None:
            cur = new(title if not slides else slides[-1]["title"] + " (계속)")
        if b.kind == "heading":
            cur["items"].append(("sub", plain(b.text), 0))
        elif b.kind == "para":
            if b.text.startswith(("노트:", "Note:", "발표자 노트:")):
                cur["notes"].append(plain(b.text.split(":", 1)[1]))
            else:
                cur["items"].append(("text", b.text, 0))
        elif b.kind == "bullet":
            cur["items"].append(("text", b.text, b.level))
        elif b.kind == "quote":
            cur["notes"].append(plain(b.text))
        elif b.kind == "table":
            cur["items"].append(("table", b.rows, 0))
        elif b.kind == "code":
            cur["items"].append(("text", b.text, 0))

    # 너무 긴 슬라이드는 (계속) 슬라이드로 나눈다
    out: list[dict] = []
    for s in slides:
        texts = [i for i in s["items"] if i[0] != "table"]
        tables = [i for i in s["items"] if i[0] == "table"]
        chunks = [texts[k:k + _MAX_LINES] for k in range(0, len(texts), _MAX_LINES)] or [[]]
        for n, chunk in enumerate(chunks):
            out.append({
                "title": s["title"] + (" (계속)" if n else ""),
                "items": chunk,
                "notes": s["notes"] if n == 0 else [],
                "table": tables[0][1] if tables and n == len(chunks) - 1 else None,
            })
        for extra in tables[1:]:
            out.append({"title": s["title"] + " (표)", "items": [], "notes": [], "table": extra[1]})
    return title, subtitle, out


def _ppt_run(paragraph, text: str, size: int, color=(0x26, 0x26, 0x26), bold_default=False) -> None:
    for chunk, bold in split_bold(text) or [("", False)]:
        run = paragraph.add_run()
        run.text = chunk
        run.font.size = PptPt(size)
        run.font.name = FONT
        run.font.bold = bold or bold_default
        run.font.color.rgb = PptColor(*color)


def write_pptx(blocks: list[Block], path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    W, H = prs.slide_width, prs.slide_height
    title, subtitle, slides = _slides_from_blocks(blocks)

    # 표지
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PptColor(*ACCENT)
    bg.line.fill.background()
    box = s.shapes.add_textbox(Inches(0.9), Inches(2.6), W - Inches(1.8), Inches(1.6))
    box.text_frame.word_wrap = True
    _ppt_run(box.text_frame.paragraphs[0], title, 40, (0xFF, 0xFF, 0xFF), True)
    if subtitle:
        p = box.text_frame.add_paragraph()
        _ppt_run(p, subtitle, 20, (0xDD, 0xE6, 0xF0))

    for num, data in enumerate(slides, start=2):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H)
        bar.fill.solid()
        bar.fill.fore_color.rgb = PptColor(*ACCENT)
        bar.line.fill.background()
        tbox = s.shapes.add_textbox(Inches(0.6), Inches(0.35), W - Inches(1.2), Inches(0.9))
        tbox.text_frame.word_wrap = True
        _ppt_run(tbox.text_frame.paragraphs[0], data["title"], 30, ACCENT, True)

        top = Inches(1.45)
        body_h = H - top - Inches(0.6)
        if data["items"]:
            has_table = data["table"] is not None
            h = body_h // 2 if has_table else body_h
            box = s.shapes.add_textbox(Inches(0.7), top, W - Inches(1.4), h)
            tf = box.text_frame
            tf.word_wrap = True
            size = 20 if len(data["items"]) <= 6 else 17
            for k, (kind, text, level) in enumerate(data["items"]):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.level = min(level, 4)
                p.space_after = PptPt(6)
                if kind == "sub":
                    _ppt_run(p, text, size + 2, ACCENT, True)
                else:
                    _ppt_run(p, ("• " if level == 0 else "– ") + text, size - 2 * min(level, 2))
            top = top + h + Inches(0.1)
            body_h = H - top - Inches(0.5)
        if data["table"]:
            rows = data["table"]
            row_h = Emu(min(int(body_h / len(rows)), int(Inches(0.5))))
            shape = s.shapes.add_table(len(rows), len(rows[0]), Inches(0.7), top,
                                       W - Inches(1.4), Emu(row_h * len(rows)))
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = shape.table.cell(r, c)
                    cell.text = ""
                    _ppt_run(cell.text_frame.paragraphs[0], val, 14 if len(rows) < 8 else 11,
                             (0xFF, 0xFF, 0xFF) if r == 0 else (0x26, 0x26, 0x26), r == 0)
        foot = s.shapes.add_textbox(W - Inches(1.2), H - Inches(0.5), Inches(0.9), Inches(0.35))
        _ppt_run(foot.text_frame.paragraphs[0], str(num), 11, (0x99, 0x99, 0x99))
        if data["notes"]:
            s.notes_slide.notes_text_frame.text = "\n".join(data["notes"])
    prs.save(path)


# ---------------------------------------------------------------- XLSX

_NUM_RE = re.compile(r"^-?[\d,]+(\.\d+)?$")


def _cell_value(text: str):
    t = plain(text).strip()
    if _NUM_RE.match(t) and not (len(t) > 1 and t.startswith("0") and "." not in t):
        try:
            v = float(t.replace(",", ""))
            return int(v) if v.is_integer() else v
        except ValueError:
            pass
    return t


def write_xlsx(blocks: list[Block], path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E79")
    used: set[str] = set()
    last_heading = "표"
    notes: list[str] = []
    has_text = False

    def sheet_name(base: str) -> str:
        base = re.sub(r"[\[\]:*?/\\]", " ", plain(base)).strip()[:28] or "표"
        name, n = base, 2
        while name in used:
            name = f"{base[:25]}_{n}"
            n += 1
        used.add(name)
        return name

    for b in blocks:
        if b.kind == "heading":
            last_heading = b.text
            notes.append(plain(b.text))
        elif b.kind == "table" and b.rows:
            ws = wb.create_sheet(sheet_name(last_heading))
            for r, row in enumerate(b.rows, start=1):
                for c, val in enumerate(row, start=1):
                    cell = ws.cell(row=r, column=c, value=plain(val) if r == 1 else _cell_value(val))
                    cell.border = border
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    cell.font = Font(name=FONT, size=10, bold=r == 1, color="FFFFFF" if r == 1 else "000000")
                    if r == 1:
                        cell.fill = head_fill
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    elif isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0" if isinstance(cell.value, int) else "#,##0.00"
            for c in range(1, len(b.rows[0]) + 1):
                width = max(len(str(ws.cell(row=r, column=c).value or "")) for r in range(1, len(b.rows) + 1))
                ws.column_dimensions[get_column_letter(c)].width = min(max(10, width * 1.8 + 2), 60)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
        elif b.kind in ("para", "bullet", "quote"):
            has_text = True
            notes.append(("  " * b.level + "• " if b.kind == "bullet" else "") + plain(b.text))

    if not wb.sheetnames or has_text:
        ws = wb.create_sheet("메모" if wb.sheetnames else "내용")
        for r, line in enumerate(notes, start=1):
            ws.cell(row=r, column=1, value=line).font = Font(name=FONT, size=10)
        ws.column_dimensions["A"].width = 100
    wb.save(path)


# ---------------------------------------------------------------- 검증 (§7)

def validate(path: Path) -> list[str]:
    ext = path.suffix.lower().lstrip(".")
    try:
        if ext == "hwpx":
            return hwpx_writer.validate_hwpx(path)
        if ext == "docx":
            Document(path)
        elif ext == "pptx":
            Presentation(path)
        elif ext == "xlsx":
            wb = load_workbook(path)
            errs = [
                f"{ws.title}!{c.coordinate}: {c.value}"
                for ws in wb.worksheets for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("#") and c.value.endswith(("!", "?", "A"))
            ]
            return [f"수식 오류 셀 {e}" for e in errs]
    except Exception as e:  # noqa: BLE001 - 어떤 손상이든 검증 실패로 보고
        return [f"파일 열기 실패: {e}"]
    return []


# ---------------------------------------------------------------- 진입점

WRITERS = {"hwpx": None, "docx": write_docx, "pptx": write_pptx, "xlsx": write_xlsx}


def render(markdown: str, fmt: str, out_dir: Path, doc_type: str = "", org: str = "") -> dict:
    fmt = fmt.lower()
    if fmt not in FORMATS:
        raise ValueError(f"지원하지 않는 형식: {fmt}")
    body, meta = extract_meta(markdown)
    blocks = parse(body)
    if not blocks:
        raise ValueError("파일로 만들 내용이 없습니다.")
    doc_type = doc_type or meta.get("doc") or title_of(blocks)
    org = org or meta.get("org", "")
    path = build_filename(out_dir, doc_type, fmt, org)
    if fmt == "hwpx":
        hwpx_writer.write_hwpx(blocks, path, title=title_of(blocks))
    else:
        WRITERS[fmt](blocks, path)
    errors = validate(path)
    return {"name": path.name, "path": str(path), "format": fmt, "valid": not errors, "errors": errors}
