"""렌더러·추출기·서버 API 테스트. 실행: python -m pytest tests -q"""

from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

SAMPLE = """[전제: 교원 대상, 4시간, PPT]

# 생성형 AI 수업 활용 연수
여주교육지원청 교원 직무연수

## 1. 연수 개요
- **대상**: 초·중등 교원 30명
- 시간: 4시간(2차시 × 2)
  - 1부: 이론
  - 2부: 실습
노트: 인사 후 연수 목표를 먼저 안내한다.

## 2. 차시별 계획
| 차시 | 주제 | 시간 | 비고 |
|---|---|---|---|
| 1 | 생성형 AI 이해 | 60 | 강의 |
| 2 | 프롬프트 실습 | 1,200 | **실습** |

### 실습 산출물
1. 수업 지도안 1건
2. 평가 루브릭

> 정책 정보는 [기준: 2026-09 / 확인 필요]

```
프롬프트 예시: 너는 초등 교사야
```

<!--meta {"org": "여주교육지원청", "doc": "강의자료"}-->
"""


@pytest.fixture()
def env(tmp_path, monkeypatch):
    from agent import config

    monkeypatch.setattr(config, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(config, "SESSION_DIR", tmp_path / "sessions")
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path / "out"))
    import agent.store as store

    monkeypatch.setattr(store, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(store, "SESSION_DIR", tmp_path / "sessions")
    config.ensure_dirs()
    return tmp_path


# ---------------------------------------------------------------- 마크다운

def test_parse_blocks():
    from agent.markdown import extract_meta, parse

    body, meta = extract_meta(SAMPLE)
    assert meta == {"org": "여주교육지원청", "doc": "강의자료"}
    assert "meta" not in body
    kinds = [b.kind for b in parse(body)]
    for k in ("heading", "para", "bullet", "table", "quote", "code"):
        assert k in kinds
    table = next(b for b in parse(body) if b.kind == "table")
    assert table.rows[0] == ["차시", "주제", "시간", "비고"]
    nested = [b for b in parse(body) if b.kind == "bullet" and b.level == 1]
    assert len(nested) == 2


# ---------------------------------------------------------------- 렌더러

@pytest.mark.parametrize("fmt", ["hwpx", "docx", "pptx", "xlsx"])
def test_render_all_formats_validate(env, fmt):
    from agent import config, extract, render

    info = render.render(SAMPLE, fmt, config.output_dir())
    assert info["valid"], info["errors"]
    assert info["name"].startswith("여주교육지원청_강의자료_")
    path = Path(info["path"])
    # 만든 파일을 다시 읽어 핵심 내용이 들어갔는지 확인
    text = extract.analyze(path)["text"]
    assert "차시별 계획" in text or fmt == "xlsx"
    assert "프롬프트 실습" in text


def test_filename_versioning(env):
    from agent import config, render

    a = render.render(SAMPLE, "docx", config.output_dir())
    b = render.render(SAMPLE, "docx", config.output_dir())
    c = render.render(SAMPLE, "docx", config.output_dir(), org="", doc_type="견적서")
    assert b["name"] == a["name"].replace(".docx", "_v2.docx")
    assert "견적서" in c["name"]


def test_hwpx_structure(env):
    from agent import config, render

    info = render.render(SAMPLE, "hwpx", config.output_dir())
    with zipfile.ZipFile(info["path"]) as zf:
        assert zf.namelist()[0] == "mimetype"
        assert zf.getinfo("mimetype").compress_type == zipfile.ZIP_STORED
        section = zf.read("Contents/section0.xml").decode()
        hpf = zf.read("Contents/content.hpf").decode()
    assert "생성형 AI 수업 활용 연수" in section
    assert "□ " in section and "○ " in section
    assert 'rowCnt="3" colCnt="4"' in section
    assert "<opf:title>생성형 AI 수업 활용 연수</opf:title>" in hpf
    ids = [int(x) for x in __import__("re").findall(r'<hp:p id="(\d+)"', section)]
    assert len(ids) == len(set(ids)), "문단 id 중복"


def test_hwpx_escapes_special_chars(env):
    from agent import config, render

    info = render.render("# A & B <테스트>\n\n- 5 < 6 & \"따옴표\"", "hwpx", config.output_dir())
    assert info["valid"], info["errors"]


def test_pptx_slides_and_notes(env):
    from pptx import Presentation

    from agent import config, render

    info = render.render(SAMPLE, "pptx", config.output_dir())
    prs = Presentation(info["path"])
    titles = [s.shapes[1].text_frame.text if len(s.shapes) > 1 else "" for s in prs.slides]
    assert len(prs.slides) >= 3
    assert any("연수 개요" in t for t in titles)
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
    assert any("연수 목표" in n for n in notes)
    assert any(sh.has_table for s in prs.slides for sh in s.shapes)


def test_pptx_long_slide_splits(env):
    from pptx import Presentation

    from agent import config, render

    md = "# 긴 자료\n\n## 목록\n" + "\n".join(f"- 항목 {i}" for i in range(20))
    prs = Presentation(render.render(md, "pptx", config.output_dir())["path"])
    assert len(prs.slides) == 1 + 3  # 표지 + 9/9/2


def test_xlsx_numbers_and_sheets(env):
    from openpyxl import load_workbook

    from agent import config, render

    info = render.render(SAMPLE, "xlsx", config.output_dir())
    wb = load_workbook(info["path"])
    ws = wb["2. 차시별 계획"]
    assert ws["C3"].value == 1200 and isinstance(ws["C3"].value, int)
    assert ws["D3"].value == "실습"
    assert ws.freeze_panes == "A2"


def test_xlsx_keeps_leading_zero_text(env):
    from openpyxl import load_workbook

    from agent import config, render

    md = "## 명단\n| 이름 | 연락처 |\n|---|---|\n| 홍길동 | 01012345678 |"
    wb = load_workbook(render.render(md, "xlsx", config.output_dir())["path"])
    assert wb["명단"]["B2"].value == "01012345678"
    assert wb.sheetnames == ["명단"]


def test_empty_markdown_rejected(env):
    from agent import config, render

    with pytest.raises(ValueError):
        render.render("   <!--meta {}-->  ", "docx", config.output_dir())


# ---------------------------------------------------------------- 추출기

def test_extract_unsupported_hwp(env):
    from agent import extract

    p = env / "old.hwp"
    p.write_bytes(b"\xd0\xcf\x11\xe0")
    info = extract.analyze(p)
    assert info["status"] == "manual"
    blocks = extract.to_blocks(p, "old.hwp", info)
    assert "수동 확인 필요" in blocks[0]["text"]


def test_extract_csv_cp949(env):
    from agent import extract

    p = env / "list.csv"
    p.write_bytes("이름,기관\n김교사,여주\n".encode("cp949"))
    assert "김교사" in extract.analyze(p)["text"]


def test_extract_pdf_as_document_block(env):
    import pymupdf

    from agent import extract

    p = env / "a.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Hello PDF")
    doc.save(p)
    info = extract.analyze(p)
    assert info["kind"] == "pdf" and "Hello PDF" in info["text"]
    assert extract.to_blocks(p, "a.pdf", info)[0]["type"] == "document"


# ---------------------------------------------------------------- LLM 루프 (SDK 스트림 모의)

class FakeStream:
    def __init__(self, events, final):
        self.events, self.final = events, final

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def __iter__(self):
        return iter(self.events)

    def get_final_message(self):
        return self.final


def _ev(kind, **kw):
    return SimpleNamespace(type=kind, **kw)


def _text_events(*chunks):
    evs = [_ev("content_block_start", content_block=SimpleNamespace(type="thinking")),
           _ev("content_block_start", content_block=SimpleNamespace(type="text"))]
    evs += [_ev("content_block_delta", delta=SimpleNamespace(type="text_delta", text=c)) for c in chunks]
    return evs


def _fake_client(monkeypatch, streams, calls):
    import agent.llm as llm

    class Messages:
        def stream(self, **kw):
            calls.append(kw)
            return streams.pop(0)

    fake = SimpleNamespace(beta=SimpleNamespace(messages=Messages()))
    monkeypatch.setattr(llm.anthropic, "Anthropic", lambda: fake)


def test_llm_run_streams_and_sets_params(env, monkeypatch):
    from agent import llm

    calls: list = []
    final = SimpleNamespace(stop_reason="end_turn", content=[])
    _fake_client(monkeypatch, [FakeStream(_text_events("안녕", "하세요"), final)], calls)
    events = list(llm.run([{"role": "user", "text": "hi", "files": []}], "gov", "claude-opus-5", "high", True))
    assert [e["type"] for e in events][-1] == "done"
    assert events[-1]["text"] == "안녕하세요"
    kw = calls[0]
    assert kw["thinking"] == {"type": "adaptive"}
    assert kw["output_config"] == {"effort": "high"}
    assert kw["fallbacks"] == "default" and kw["betas"] == [llm.FALLBACK_BETA]
    assert kw["tools"][0]["type"] == "web_search_20260209"
    assert "정부지원사업" in kw["system"][1]["text"] and "HWPX" in kw["system"][1]["text"]


def test_llm_run_sonnet_no_fallback_and_pause_turn(env, monkeypatch):
    from agent import llm

    calls: list = []
    paused = SimpleNamespace(stop_reason="pause_turn", content=[{"type": "text", "text": "A"}])
    final = SimpleNamespace(stop_reason="end_turn", content=[])
    _fake_client(monkeypatch, [FakeStream(_text_events("A"), paused), FakeStream(_text_events("B"), final)], calls)
    events = list(llm.run([{"role": "user", "text": "hi", "files": []}], "auto", "claude-sonnet-5", "low", False))
    assert "fallbacks" not in calls[0] and "tools" not in calls[0]
    assert calls[1]["messages"][-1]["role"] == "assistant"
    assert events[-1]["text"] == "A\n\nB"


def test_llm_refusal_reports_error(env, monkeypatch):
    from agent import llm

    calls: list = []
    _fake_client(monkeypatch, [FakeStream([], SimpleNamespace(stop_reason="refusal", content=[]))], calls)
    events = list(llm.run([{"role": "user", "text": "x", "files": []}], "auto", "claude-opus-5", "low", False))
    assert events[-1]["type"] == "error"


# ---------------------------------------------------------------- 서버 API

def test_server_end_to_end(env, monkeypatch):
    from fastapi.testclient import TestClient

    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test")
    import server

    def fake_run(history, mode, model, effort, web, fmt=""):
        assert history[-1]["text"] == "연수 PPT 만들어줘"
        assert fmt == "pptx"
        yield {"type": "text", "text": SAMPLE}
        yield {"type": "done", "text": SAMPLE, "stop_reason": "end_turn"}

    monkeypatch.setattr(server.llm, "run", fake_run)
    c = TestClient(server.app)

    st = c.get("/api/status").json()
    assert st["has_key"] and len(st["modes"]) == 7
    s = c.post("/api/sessions", json={"mode": "lecture"}).json()

    up = c.post("/api/upload", files=[("files", ("명단.csv", "이름\n김\n".encode(), "text/csv")),
                                      ("files", ("구형.hwp", b"xx", "application/octet-stream"))]).json()
    assert up[0]["status"] == "ok" and up[1]["status"] == "manual"

    r = c.post("/api/chat", json={"session_id": s["id"], "message": "연수 PPT 만들어줘",
                                  "file_ids": [up[0]["id"]], "mode": "lecture"})
    events = [json.loads(line[6:]) for line in r.text.split("\n\n") if line.startswith("data: ")]
    types = [e["type"] for e in events]
    assert "saved" in types and types[-1] == "end"
    file_ev = next(e for e in events if e["type"] == "file")
    assert file_ev["format"] == "pptx" and file_ev["valid"]

    r2 = c.post("/api/render", json={"session_id": s["id"], "index": 1, "format": "hwpx"}).json()
    assert r2["valid"] and r2["name"].endswith(".hwpx")

    saved = c.get(f"/api/sessions/{s['id']}").json()
    assert saved["title"] == "연수 PPT 만들어줘" and len(saved["outputs"]) == 2
    assert saved["messages"][0]["file_info"][0]["name"] == "명단.csv"

    names = [o["name"] for o in c.get("/api/outputs").json()]
    assert file_ev["name"] in names
    assert c.get(f"/api/download/{file_ev['name']}").status_code == 200
    assert c.get("/api/download/..%2F..%2Fserver.py").status_code == 404
    assert c.get("/api/sessions/../../etc").status_code == 404
    assert c.delete(f"/api/sessions/{s['id']}").json()["deleted"]


def test_server_failed_chat_rolls_back(env, monkeypatch):
    from fastapi.testclient import TestClient

    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test")
    import server

    monkeypatch.setattr(server.llm, "run", lambda *a, **k: iter([{"type": "error", "text": "API 키 오류"}]))
    c = TestClient(server.app)
    s = c.post("/api/sessions", json={}).json()
    r = c.post("/api/chat", json={"session_id": s["id"], "message": "hi"})
    assert "API 키 오류" in r.text
    assert c.get(f"/api/sessions/{s['id']}").json()["messages"] == []


def test_settings_rejects_bad_key(env, monkeypatch):
    from fastapi.testclient import TestClient

    import server

    monkeypatch.setattr(server.config, "ENV_PATH", env / ".env")
    c = TestClient(server.app)
    assert c.post("/api/settings", json={"api_key": "wrong"}).status_code == 400
    ok = c.post("/api/settings", json={"api_key": "sk-ant-abc", "model": "claude-sonnet-5"})
    assert ok.status_code == 200 and ok.json()["default_model"] == "claude-sonnet-5"
    assert "ANTHROPIC_API_KEY=sk-ant-abc" in (env / ".env").read_text(encoding="utf-8")
